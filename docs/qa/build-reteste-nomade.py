from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"C:\Users\Teste\Documentos 2.0\Verificação 2.0\_DESENVOLVIMENTOS\Plataforma allka 2026\allka-plataforma\allka-2026")
OUT = ROOT / "entregas" / "qa-usuarios" / "atualizacao-nomade"
OUT.mkdir(parents=True, exist_ok=True)

tests = [
    ("NO-R01", "Login e sessão", "Acesse https://allka.store com a conta Nômade de QA.", "O login abre o portal do Nômade e a sessão permanece ativa ao navegar."),
    ("NO-R02", "Container padrão", "Na tela inicial, confira cabeçalho, menu lateral e conteúdo central.", "Nenhuma área fica cortada ou sobreposta. O conteúdo usa o mesmo container central da plataforma."),
    ("NO-R03", "Tutorial no primeiro acesso", "Caso a oferta apareça, clique em Começar.", "O primeiro passo do tutorial abre e o balão destaca o elemento correto."),
    ("NO-R04", "Não quero ver novamente", "Na oferta do tutorial, escolha Não quero ver mais este tutorial e recarregue a página.", "A oferta não reaparece sozinha após F5. O tutorial continua disponível manualmente pela Ajuda."),
    ("NO-R05", "Central de Ajuda", "Abra o ícone de Ajuda e localize o tutorial do portal Nômade.", "O ícone de Ajuda é diferente do ícone de Sugestões. É possível iniciar, continuar ou refazer o tour."),
    ("NO-R06", "Retomada do tutorial", "Inicie um tour, avance dois passos, feche e recarregue a página.", "O tour oferece continuidade no mesmo ponto; não volta silenciosamente ao início."),
    ("NO-R07", "Allkademy", "Abra Allkademy pelo menu do Nômade.", "A tela abre com cabeçalho, menu lateral e container central padronizados."),
    ("NO-R08", "Habilitações", "Abra Habilitações e confirme as áreas disponíveis para o perfil de QA.", "A tela abre sem erro e explica claramente as habilitações. Registre se as áreas necessárias para os testes foram configuradas."),
    ("NO-R09", "Minhas Tarefas em lista", "Abra Minhas Tarefas e selecione Lista.", "As tarefas atribuídas aparecem como cards legíveis, com prazo, etapa e estado."),
    ("NO-R10", "Minhas Tarefas em kanban", "Na mesma tela, selecione Kanban e volte para Lista.", "As duas visualizações funcionam e mostram a mesma situação das tarefas."),
    ("NO-R11", "Detalhe da tarefa", "Clique em uma tarefa ou etapa atribuída.", "O detalhe abre dentro do container central, no tamanho da área principal, com fechar e fixar na bandeja de telas."),
    ("NO-R12", "Etapa concluída", "Em uma tarefa marcada [TESTE USUÁRIO], conclua apenas a etapa liberada e abra a aba Concluídas.", "A tarefa/etapa concluída permanece visível na aba Concluídas, sem desaparecer."),
    ("NO-R13", "Sequência automática", "Em uma tarefa [TESTE USUÁRIO] com duas etapas sequenciais, conclua a primeira.", "A próxima etapa é liberada automaticamente. Antes disso, ela aparece bloqueada e explica o motivo."),
    ("NO-R14", "Histórico", "Abra Histórico após uma ação de teste permitida.", "A ação aparece com data e descrição coerentes."),
    ("NO-R15", "Tarefas Disponíveis", "Abra Tarefas Disponíveis. Se existir uma etapa [TESTE USUÁRIO] compatível com uma habilitação, assuma-a.", "A busca filtra normalmente. Após assumir uma etapa de teste, ela passa a aparecer em Minhas Tarefas."),
    ("NO-R16", "Permissões do Nômade", "Percorra o menu e tente localizar catálogo de contratação.", "O Nômade não vê catálogo nem ações de compra. Só aparecem recursos próprios do perfil."),
]

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, val in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')

def style_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    for idx, cell in enumerate(table.rows[0].cells):
        set_cell_shading(cell, '14213D')
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
    for row_i, row in enumerate(table.rows[1:], start=1):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(9)
        if row_i % 2 == 0:
            for cell in row.cells:
                set_cell_shading(cell, 'F4F7FB')

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.6)
section.left_margin = Inches(0.7)
section.right_margin = Inches(0.7)
styles = doc.styles
styles['Normal'].font.name = 'Arial'
styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
styles['Normal'].font.size = Pt(10)
for name, size in [('Title', 22), ('Heading 1', 15), ('Heading 2', 12)]:
    styles[name].font.name = 'Arial'
    styles[name]._element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
    styles[name]._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
    styles[name].font.size = Pt(size)
    styles[name].font.color.rgb = RGBColor(0, 0, 0)

title = doc.add_paragraph(style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.add_run('Roteiro de Reteste Nômade')
subtitle = doc.add_paragraph()
subtitle.add_run('Atualização de QA - setembro de 2026').italic = True
intro = doc.add_paragraph()
intro.add_run('Objetivo. ').bold = True
intro.add_run('Este roteiro verifica as correções recentes do portal Nômade antes de liberar o fluxo para o uso normal. Execute apenas com a conta de QA indicada e registre cada resultado na planilha correspondente.')

doc.add_heading('Acesso de teste', level=1)
t = doc.add_table(rows=1, cols=2)
t.rows[0].cells[0].text = 'Campo'
t.rows[0].cells[1].text = 'Informação'
for a, b in [
    ('Portal', 'https://allka.store'),
    ('Conta', 'nomad@allka.com.vc'),
    ('Perfil', 'Nômade de QA, ativo, nível Bronze'),
    ('Senha', 'Enviar separadamente por canal interno seguro'),
]:
    row = t.add_row().cells
    row[0].text = a
    row[1].text = b
style_table(t)

doc.add_heading('Antes de começar', level=1)
for text in [
    'Use apenas registros marcados [TESTE USUÁRIO] para concluir etapas ou assumir tarefas.',
    'Os testes NO-R12 a NO-R15 precisam de uma tarefa de QA atribuída e de uma tarefa disponível compatível. Se elas não existirem, marque Não consegui testar e informe a coordenação.',
    'Não exclua, cancele ou altere dados reais. Não publique produtos, não faça pagamentos e não altere contas.',
    'Ao encontrar algo errado, tire uma captura de tela, informe o código do teste e descreva o que esperava acontecer.',
]:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(text)

doc.add_heading('Roteiro de testes', level=1)
for code, name, steps, expected in tests:
    doc.add_heading(f'{code} {name}', level=2)
    p = doc.add_paragraph()
    p.add_run('Como testar: ').bold = True
    p.add_run(steps)
    p = doc.add_paragraph()
    p.add_run('Deve acontecer: ').bold = True
    p.add_run(expected)

doc.add_heading('Como registrar', level=1)
doc.add_paragraph('Na planilha, selecione Funcionou, Não funcionou, Não se aplica ou Não consegui testar. Em caso de problema, descreva o que ocorreu, inclua a URL da tela e anexe a captura quando possível.')
doc.add_heading('Pendências conhecidas fora deste reteste', level=1)
doc.add_paragraph('Em telas estreitas, os ícones flutuantes da lateral direita podem ficar parcialmente atrás do container principal. Esse acabamento já foi identificado e não faz parte deste reteste; registre apenas se o comportamento piorar ou impedir uma ação principal.')

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('Allka - roteiro interno de QA')

out_file = OUT / 'Roteiro_de_Reteste_Nomade.docx'
doc.save(out_file)
print(out_file)
