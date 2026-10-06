from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "documentos" / "Registro de Continuidade Produtos Allka.docx"
OUT.parent.mkdir(parents=True, exist_ok=True)

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.55)
section.bottom_margin = Inches(0.55)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
styles["Normal"].font.size = Pt(9)

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)

def add_heading(text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(8 if level == 1 else 5)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.paragraph_format.space_after = Pt(6)
run = title.add_run("Registro de Continuidade Produtos Allka")
run.font.color.rgb = RGBColor(0, 0, 0)
run.font.size = Pt(22)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(10)
p.add_run("Objetivo: ").bold = True
p.add_run("manter o contexto de trabalho do cadastro de produtos e permitir a continuidade exata da conversa entre Vinicius, Cláudio e os próximos atendimentos.")

add_heading("Estado atual", 1)
for text in [
    "Os produtos do catálogo foram deixados somente com seus títulos, sem preço, tarefas, etapas, categoria ou demais preenchimentos, para cadastro real item a item.",
    "Não houve novo deploy nesta etapa. As validações e ajustes estão sendo feitos no ambiente local.",
    "O produto e o editor permanecem em rascunho. Nada deve ser publicado automaticamente.",
]:
    doc.add_paragraph(text, style="List Bullet")

add_heading("Decisões confirmadas", 1)
rows = [
    ("Edição de tarefa", "A ação Editar tarefa só aparece depois de liberar o editor geral pelo botão Editar. Com o produto travado, ela não aparece."),
    ("Campos da tarefa", "Campos devem ter nomes claros. Executor, especialidade ou conhecimento humano, ciclo, repetição, regra de acessos e continuidade são configurados na tarefa."),
    ("Horas", "A tarefa não recebe tempo próprio. Todas as horas são informadas nas etapas. A tarefa exibe a soma automática das horas das etapas."),
    ("Preço", "O preço comercial do produto é calculado a partir das etapas e das especialidades configuradas, com taxas e margem. O cartão da tarefa exibe o custo direto estimado das etapas, sem misturar a composição comercial."),
    ("Precificação", "Especialidades, conhecimentos humanos, perfis e custos de IA são mantidos na tela global de Precificação, não dentro da tarefa."),
    ("Retorno para a tarefa", "Antes de abrir a Precificação, o rascunho em edição é salvo. Ao salvar na Precificação, o sistema retorna para a mesma tarefa, aberta no mesmo ponto."),
    ("Fechar Precificação", "O X fica dentro do cabeçalho da tela de Precificação. Fechar sem salvar retorna à tarefa e preserva o rascunho dela, sem alterar os valores globais."),
    ("Accordions", "Somente cards principais recolhem. Ao abrir um card principal, os campos internos permanecem visíveis; não deve haver um segundo nível de acordeões."),
]
for topic, decision in rows:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    lead = p.add_run(f"{topic}: ")
    lead.bold = True
    p.add_run(decision)

add_heading("Implementado nesta sessão", 1)
for text in [
    "Rótulos explícitos nos campos da edição de tarefa.",
    "Botão Editar tarefa oculto quando o produto está em modo travado.",
    "Atalho da tarefa para Precificação com salvamento de rascunho e retorno ao mesmo contexto.",
    "Tela de Precificação ampliada, com botão de fechar dentro do cabeçalho.",
    "Horas migradas visualmente para as etapas: o formulário de etapa usa Horas estimadas e converte para a base interna de cálculo.",
    "Tarefa agora apresenta total de horas e custo direto estimado das etapas quando possível; o preço comercial continua sendo calculado automaticamente no produto.",
]:
    doc.add_paragraph(text, style="List Bullet")

add_heading("Validação pendente", 1)
for text in [
    "Criar uma tarefa real e pelo menos duas etapas com horas e especialidades para validar a soma de horas, o custo direto no cartão da tarefa e o preço comercial do produto.",
    "Testar o fluxo editar tarefa, abrir Precificação, salvar um valor e confirmar o retorno automático à mesma tarefa.",
    "Testar o X da Precificação sem salvar, confirmando que volta à tarefa sem alterar os valores globais.",
    "Conferir visualmente que o cartão da tarefa chama o valor de custo estimado e que o preço comercial do produto é exibido no resumo de precificação.",
]:
    doc.add_paragraph(text, style="List Bullet")

add_heading("Última solicitação atendida", 1)
doc.add_paragraph("Remover a duração da tarefa, usar horas somente nas etapas, somar automaticamente o total de horas por tarefa e apresentar o preço estimado. Registrar o contexto da reunião e da continuidade em um documento compartilhável para que o trabalho possa ser retomado exatamente deste ponto.")

doc.save(OUT)
print(OUT)
