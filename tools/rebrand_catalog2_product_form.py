from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs' / 'Modelo Cadastro Completo de Produto Allka.docx'
OUTPUT = ROOT / 'docs' / 'Modelo Cadastro Completo de Produto Allka Plataforma.docx'
LOGO = ROOT / 'apps' / 'frontend' / 'public' / 'logo-allka-full.png'

NAVY, PURPLE, PINK, LILAC, PALE, GRID = '081A3C', '5125D8', 'D91D8E', 'EEE9FF', 'FAFAFE', 'D9D9D9'

def shade(cell_or_paragraph, fill):
    element = cell_or_paragraph._tc if hasattr(cell_or_paragraph, '_tc') else cell_or_paragraph._p
    pr = element.get_or_add_tcPr() if hasattr(cell_or_paragraph, '_tc') else element.get_or_add_pPr()
    shd = pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd'); pr.append(shd)
    shd.set(qn('w:fill'), fill)

def border_cell(cell, color=GRID):
    pr = cell._tc.get_or_add_tcPr()
    borders = pr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders'); pr.append(borders)
    for edge in ('top','left','bottom','right'):
        el = borders.find(qn(f'w:{edge}'))
        if el is None:
            el = OxmlElement(f'w:{edge}'); borders.append(el)
        el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '6'); el.set(qn('w:color'), color)

def set_margins(cell, v=100, h=130):
    pr = cell._tc.get_or_add_tcPr()
    mar = pr.first_child_found_in('w:tcMar')
    if mar is None:
        mar = OxmlElement('w:tcMar'); pr.append(mar)
    for side, value in [('top',v),('start',h),('bottom',v),('end',h)]:
        el = mar.find(qn(f'w:{side}'))
        if el is None:
            el = OxmlElement(f'w:{side}'); mar.append(el)
        el.set(qn('w:w'), str(value)); el.set(qn('w:type'), 'dxa')

def format_run(run, size=None, color=None, bold=None):
    run.font.name = 'Aptos'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    if size: run.font.size = Pt(size)
    if color: run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None: run.bold = bold

def header(section):
    h = section.header
    h.is_linked_to_previous = False
    p = h.paragraphs[0]
    p.text = ''
    p.paragraph_format.space_after = Pt(0)
    table = h.add_table(rows=1, cols=3, width=Inches(7.1))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    fills = [NAVY, PURPLE, PINK]
    widths = [2.4, 2.3, 2.4]
    for i,cell in enumerate(table.rows[0].cells):
        cell.width = Inches(widths[i]); shade(cell, fills[i]); set_margins(cell, 60, 100)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if i == 0:
            p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.add_run().add_picture(str(LOGO), width=Inches(.9))
        elif i == 1:
            p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run('CADASTRO DE PRODUTOS'); format_run(r, 7.5, 'FFFFFF', True)
        else:
            p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r = p.add_run('ALLKA  |  PLATAFORMA'); format_run(r, 7.5, 'FFFFFF', True)

def footer(section):
    f = section.footer
    p = f.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.text = ''
    r = p.add_run('ALLKA  |  FORMULÁRIO DE CADASTRO COMPLETO  |  PREENCHA ANTES DE CADASTRAR')
    format_run(r, 8, '64748B', True)

def insert_cover_banner(doc):
    first = doc.paragraphs[0]
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i,cell in enumerate(table.rows[0].cells):
        shade(cell, [NAVY,PURPLE,PINK][i]); set_margins(cell, 180, 180)
        if i == 0:
            p = cell.paragraphs[0]; p.add_run().add_picture(str(LOGO), width=Inches(1.35))
        elif i == 1:
            p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r=p.add_run('PLATAFORMA ALLKA'); format_run(r, 9, 'FFFFFF', True)
        else:
            p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r=p.add_run('PRODUTO 100% PRONTO'); format_run(r, 9, 'FFFFFF', True)
    first._p.addprevious(table._tbl)

def style_doc(doc):
    for section in doc.sections:
        section.top_margin = Inches(.72); section.bottom_margin = Inches(.65)
        section.left_margin = Inches(.7); section.right_margin = Inches(.7)
        header(section); footer(section)

    for paragraph in doc.paragraphs:
        style = paragraph.style.name if paragraph.style else ''
        if style == 'Title':
            paragraph.paragraph_format.space_before = Pt(15)
            paragraph.paragraph_format.space_after = Pt(7)
            for run in paragraph.runs: format_run(run, 25, '000000', True)
        elif style == 'Heading 1':
            shade(paragraph, LILAC)
            paragraph.paragraph_format.space_before = Pt(15)
            paragraph.paragraph_format.space_after = Pt(6)
            paragraph.paragraph_format.left_indent = Inches(.08)
            for run in paragraph.runs: format_run(run, 13, '000000', True)
        elif style == 'Heading 2':
            paragraph.paragraph_format.space_before = Pt(11)
            paragraph.paragraph_format.space_after = Pt(4)
            for run in paragraph.runs: format_run(run, 11, '000000', True)
        else:
            for run in paragraph.runs: format_run(run, 10.2)

    for table in doc.tables:
        for ri,row in enumerate(table.rows):
            for cell in row.cells:
                set_margins(cell)
                border_cell(cell)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if ri == 0 and len(table.rows) > 1 and len(row.cells) > 2:
                    shade(cell, NAVY)
                    for p in cell.paragraphs:
                        for r in p.runs: format_run(r, 8.5, 'FFFFFF', True)
                elif ri % 2 == 0:
                    shade(cell, PALE)
                else:
                    shade(cell, 'FFFFFF')
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(0)
                    for r in p.runs:
                        if not (ri == 0 and len(table.rows)>1 and len(row.cells)>2): format_run(r, 9)

def main():
    doc = Document(SOURCE)
    insert_cover_banner(doc)
    style_doc(doc)
    doc.save(OUTPUT)
    print(OUTPUT)

if __name__ == '__main__':
    main()
