from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path(__file__).resolve().parents[1] / "docs" / "Modelo Cadastro Completo de Produto Allka.docx"

NAVY = "0A1633"
PURPLE = "6124E8"
PINK = "D91D8E"
PALE = "F4F1FF"
LIGHT = "F7F9FC"
GRID = "D9D9D9"
MUTED = "5D6B82"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def borders(cell, color=GRID):
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.first_child_found_in('w:tcBorders')
    if b is None:
        b = OxmlElement('w:tcBorders')
        tcPr.append(b)
    for edge in ('top','left','bottom','right'):
        tag = qn(f'w:{edge}')
        el = b.find(tag)
        if el is None:
            el = OxmlElement(f'w:{edge}')
            b.append(el)
        el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'6'); el.set(qn('w:color'),color)

def set_cell_text(cell, text, bold=False, color=None, size=9):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.name = 'Aptos'
    r._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    r._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    if color: r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    borders(cell)

def set_width(cell, inches):
    cell.width = Inches(inches)

def cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    mar = tcPr.first_child_found_in('w:tcMar')
    if mar is None:
        mar = OxmlElement('w:tcMar'); tcPr.append(mar)
    for side, val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = mar.find(qn(f'w:{side}'))
        if node is None:
            node = OxmlElement(f'w:{side}'); mar.append(node)
        node.set(qn('w:w'), str(val)); node.set(qn('w:type'), 'dxa')

def style_table(table, header=True):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    for ri,row in enumerate(table.rows):
        for c in row.cells:
            cell_margins(c)
            borders(c)
            if ri == 0 and header:
                shade(c, NAVY)
                for r in c.paragraphs[0].runs:
                    r.font.color.rgb = RGBColor(255,255,255)
                    r.font.bold = True

def add_title(doc, text, subtitle=None):
    p = doc.add_paragraph(style='Title')
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    r.font.name = 'Aptos Display'; r.font.size = Pt(26); r.font.bold = True; r.font.color.rgb = RGBColor(0,0,0)
    if subtitle:
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_after = Pt(16)
        r2 = p2.add_run(subtitle); r2.font.size = Pt(11); r2.font.color.rgb = RGBColor.from_string(MUTED)

def heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f'Heading {level}'
    p.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text); r.font.color.rgb = RGBColor(0,0,0); r.font.name = 'Aptos Display'
    return p

def note(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(text); r.italic = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor.from_string(MUTED)

def normal(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text); r.font.size = Pt(10.5); r.font.name = 'Aptos'
    return p

def simple_form(doc, rows, widths=(2.0, 4.6)):
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for label, value in rows:
        cells = table.add_row().cells
        set_width(cells[0], widths[0]); set_width(cells[1], widths[1])
        shade(cells[0], PALE); shade(cells[1], LIGHT)
        set_cell_text(cells[0], label, bold=True, color=NAVY)
        set_cell_text(cells[1], value)
        for c in cells: cell_margins(c, 105, 115, 105, 115)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table

def matrix(doc, headers, rows, widths=None, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, color='FFFFFF', size=font_size)
        shade(table.rows[0].cells[i], NAVY)
        if widths: set_width(table.rows[0].cells[i], widths[i])
    for ridx,row in enumerate(rows):
        cells = table.add_row().cells
        for i,value in enumerate(row):
            shade(cells[i], 'FFFFFF' if ridx%2==0 else 'F5F7FB')
            set_cell_text(cells[i], value, size=font_size)
            if widths: set_width(cells[i], widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table

def line(doc, label, lines=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(label + '  '); r.bold = True; r.font.size = Pt(10)
    r2 = p.add_run('_' * 78); r2.font.color.rgb = RGBColor.from_string('98A3B5')
    for _ in range(lines-1):
        q = doc.add_paragraph('_' * 100); q.paragraph_format.space_after = Pt(3); q.runs[0].font.color.rgb = RGBColor.from_string('98A3B5')

def page_break(doc):
    doc.add_page_break()

def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Allka  |  Modelo de cadastro completo de produto  |  Preencha antes de cadastrar na plataforma')
    r.font.size = Pt(8); r.font.color.rgb = RGBColor.from_string(MUTED)

def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.6); sec.left_margin = Inches(.7); sec.right_margin = Inches(.7)
    add_footer(sec)

    styles = doc.styles
    styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(10.5)
    styles['Title'].font.color.rgb = RGBColor(0,0,0)
    styles['Heading 1'].font.color.rgb = RGBColor(0,0,0)
    styles['Heading 2'].font.color.rgb = RGBColor(0,0,0)

    add_title(doc, 'Cadastro Completo de Produto Allka', 'Formulário de preparação para cadastrar, revisar, testar e publicar um produto no catálogo da plataforma.')
    normal(doc, 'Preencha um documento por produto. Este material transforma a definição comercial e operacional em um cadastro verificável. Ao final, a pessoa responsável pelo cadastro usa cada seção no Cadastro de Produtos da Allka.')
    heading(doc, 'Como usar este formulário', 1)
    for t in [
        'Preencha todos os campos aplicáveis antes de iniciar o cadastro. Onde o produto não tiver variação, adicional, condição ou recorrência, marque “não se aplica”.',
        'Não use valores provisórios como decisão comercial. Um produto só está 100% pronto quando conteúdo, classificação, tarefas, prazo e precificação foram definidos e revisados.',
        'As perguntas de briefing pertencem às tarefas. Se duas tarefas usam o mesmo questionário, registre o mesmo nome para que ele seja reutilizado pela plataforma.',
        'O preço do especialista na plataforma é o valor/hora global da especialidade, não um campo manual exclusivo do produto. Registre-o aqui para conferência e alinhamento comercial.',
    ]:
        p=doc.add_paragraph(style='List Bullet'); p.add_run(t)
    heading(doc, 'Controle do documento', 1)
    simple_form(doc, [
        ('Produto', '____________________________________________________________'),
        ('Responsável pelo preenchimento', '____________________________________________________________'),
        ('Data de início', '____ / ____ / ________'),
        ('Revisado por', '____________________________________________________________'),
        ('Data da revisão', '____ / ____ / ________'),
        ('Situação deste formulário', '☐ Em elaboração   ☐ Pronto para cadastro   ☐ Cadastrado   ☐ Testado   ☐ Pronto para publicação'),
    ])

    page_break(doc)
    add_title(doc, 'Dados base do produto', 'Seção correspondente à criação e às Informações do produto no Cadastro de Produtos.')
    heading(doc, '1 Identificação e presença comercial', 1)
    simple_form(doc, [
        ('Nome interno do produto', '____________________________________________________________'),
        ('Slug único', '____________________________________________________________'),
        ('Título comercial', '____________________________________________________________'),
        ('Origem', '☐ Novo   ☐ Existente   ☐ Reativado'),
        ('Status inicial', '☐ Em preparação   ☐ Disponível após publicação   ☐ Temporariamente inativo   ☐ Arquivado'),
        ('Imagem principal', 'Arquivo / link / responsável: ____________________________________________'),
        ('Badge comercial opcional', '☐ Novo   ☐ Lançamento   ☐ Promoção   ☐ Destaque   ☐ Nenhum'),
        ('Texto e validade da promoção', '____________________________________________________________'),
    ])
    heading(doc, '2 Texto para o catálogo e detalhes', 1)
    line(doc, 'Descrição curta', 3)
    line(doc, 'Descrição completa', 7)
    line(doc, 'O que está incluído', 5)
    line(doc, 'Principais destaques', 4)
    line(doc, 'Informações adicionais e limites de escopo', 4)
    line(doc, 'Resumo da mudança para o histórico da versão', 2)

    page_break(doc)
    add_title(doc, 'Classificação e oferta', 'Dados que determinam em quais filtros e áreas o produto aparece.')
    heading(doc, '3 Classificação', 1)
    simple_form(doc, [
        ('Pilar', '____________________________________________________________'),
        ('Categoria', '____________________________________________________________'),
        ('Classificações 4F', '☐ __________   ☐ __________   ☐ __________   ☐ __________'),
        ('Produto recorrente mensal?', '☐ Sim, a entrega se repete a cada mês   ☐ Não, é entrega avulsa'),
        ('Observação de recorrência', '____________________________________________________________'),
    ])
    heading(doc, '4 Modalidades de contratação por período', 1)
    note(doc, 'Na operação atual, só a modalidade mensal pode ficar ativa. Trimestral, semestral e anual permanecem registradas para ativação futura. Preencha desconto somente quando houver decisão comercial explícita.')
    matrix(doc, ['Período', 'Meses', 'Desconto', 'Ativo?', 'Observação'], [
        ['Mensal', '1', '______ %', '☐ Sim  ☐ Não', '________________________________'],
        ['Trimestral', '3', '______ %', 'Reservado para o futuro', '________________________________'],
        ['Semestral', '6', '______ %', 'Reservado para o futuro', '________________________________'],
        ['Anual', '12', '______ %', 'Reservado para o futuro', '________________________________'],
    ], [1.1,.7,1.0,1.7,2.2])
    heading(doc, '5 Variações obrigatórias', 1)
    note(doc, 'Variação é uma escolha obrigatória da contratação, como formato, volume, tamanho ou modalidade. Adicionais opcionais ficam na seção seguinte.')
    matrix(doc, ['Chave', 'Nome da variação', 'Obrigatória?', 'Seleção', 'Observações'], [
        ['________________', '________________________', '☐ Sim ☐ Não', '☐ Única ☐ Múltipla', '____________________________'],
        ['________________', '________________________', '☐ Sim ☐ Não', '☐ Única ☐ Múltipla', '____________________________'],
        ['________________', '________________________', '☐ Sim ☐ Não', '☐ Única ☐ Múltipla', '____________________________'],
    ], [1.2,1.7,1.1,1.1,1.5])
    matrix(doc, ['Variação', 'Chave da opção', 'Rótulo para o cliente', 'Padrão?', 'Efeito de prazo, preço, tarefa ou informação'], [
        ['________________', '________________', '____________________________', '☐ Sim ☐ Não', '________________________________________'],
        ['________________', '________________', '____________________________', '☐ Sim ☐ Não', '________________________________________'],
        ['________________', '________________', '____________________________', '☐ Sim ☐ Não', '________________________________________'],
        ['________________', '________________', '____________________________', '☐ Sim ☐ Não', '________________________________________'],
    ], [1.2,1.2,1.7,.8,2.0], 8)

    page_break(doc)
    add_title(doc, 'Adicionais e regras de contratação', 'Escolhas opcionais e regras que alteram preço, prazo, tarefas ou informações pedidas.')
    heading(doc, '6 Adicionais opcionais', 1)
    matrix(doc, ['Chave', 'Nome', 'Descrição para o cliente', 'Custo direto', 'Padrão?', 'Ativo?', 'Vínculo tarefa ou etapa'], [
        ['____________', '________________', '________________________', 'R$ ________', '☐ Sim ☐ Não', '☐ Sim ☐ Não', '________________________'],
        ['____________', '________________', '________________________', 'R$ ________', '☐ Sim ☐ Não', '☐ Sim ☐ Não', '________________________'],
        ['____________', '________________', '________________________', 'R$ ________', '☐ Sim ☐ Não', '☐ Sim ☐ Não', '________________________'],
    ], [.9,1.2,1.6,1.0,.8,.7,1.3], 8)
    matrix(doc, ['Adicional', 'Efeito', 'Valor do efeito', 'Detalhe'], [
        ['________________', '☐ + dias ☐ + valor ☐ + % ☐ + tarefa ☐ + etapa ☐ pedir informação ☐ entregável', '________________', '____________________________'],
        ['________________', '☐ + dias ☐ + valor ☐ + % ☐ + tarefa ☐ + etapa ☐ pedir informação ☐ entregável', '________________', '____________________________'],
    ], [1.2,3.2,1.2,1.4], 8)
    heading(doc, '7 Condições', 1)
    note(doc, 'Condição é uma regra tipada: gatilho → comparação → efeito. Use chaves existentes de variação, adicional, tarefa ou atributo.')
    matrix(doc, ['Chave e nome', 'Gatilho', 'Referência', 'Operador e comparação', 'Efeito e valor', 'Ativa?'], [
        ['________________________', '☐ Opção ☐ Adicional ☐ Quantidade ☐ Resposta ☐ Atributo', '____________', '____________________', '____________________', '☐ Sim ☐ Não'],
        ['________________________', '☐ Opção ☐ Adicional ☐ Quantidade ☐ Resposta ☐ Atributo', '____________', '____________________', '____________________', '☐ Sim ☐ Não'],
        ['________________________', '☐ Opção ☐ Adicional ☐ Quantidade ☐ Resposta ☐ Atributo', '____________', '____________________', '____________________', '☐ Sim ☐ Não'],
    ], [1.4,1.7,.8,1.4,1.3,.7], 7.5)

    page_break(doc)
    add_title(doc, 'Entrega tarefas etapas e briefing', 'Uma linha por tarefa do produto. Duplique esta página se houver mais tarefas.')
    heading(doc, '8 Tarefas do produto', 1)
    matrix(doc, ['Ordem', 'Chave da tarefa', 'Nome da tarefa', 'Especialidade', 'Modo', 'Minutos', 'Revisão?', 'Aprovação cliente?', 'Condicional?'], [
        ['____', '________________', '________________________', '________________', '☐ Humano ☐ IA ☐ Híbrido', '______', '☐ Sim ☐ Não', '☐ Sim ☐ Não', '☐ Sim ☐ Não'],
        ['____', '________________', '________________________', '________________', '☐ Humano ☐ IA ☐ Híbrido', '______', '☐ Sim ☐ Não', '☐ Sim ☐ Não', '☐ Sim ☐ Não'],
        ['____', '________________', '________________________', '________________', '☐ Humano ☐ IA ☐ Híbrido', '______', '☐ Sim ☐ Não', '☐ Sim ☐ Não', '☐ Sim ☐ Não'],
        ['____', '________________', '________________________', '________________', '☐ Humano ☐ IA ☐ Híbrido', '______', '☐ Sim ☐ Não', '☐ Sim ☐ Não', '☐ Sim ☐ Não'],
    ], [.45,1.0,1.4,1.1,1.1,.55,.7,.9,.7], 7.2)
    line(doc, 'Descrição e objetivo das tarefas (relacione pela chave)', 5)
    heading(doc, '9 Etapas de cada tarefa', 1)
    matrix(doc, ['Tarefa', 'Chave da etapa', 'Nome da etapa', 'Descrição', 'Minutos', 'Condicional?'], [
        ['____________', '____________', '__________________', '________________________', '______', '☐ Sim ☐ Não'],
        ['____________', '____________', '__________________', '________________________', '______', '☐ Sim ☐ Não'],
        ['____________', '____________', '__________________', '________________________', '______', '☐ Sim ☐ Não'],
        ['____________', '____________', '__________________', '________________________', '______', '☐ Sim ☐ Não'],
    ], [1.0,1.0,1.3,2.1,.7,1.0], 8)
    heading(doc, '10 Dependências entre tarefas', 1)
    matrix(doc, ['Tarefa que depende', 'Depende desta tarefa', 'Motivo'], [
        ['________________________', '________________________', '________________________________________'],
        ['________________________', '________________________', '________________________________________'],
    ], [2.0,2.0,2.8])

    page_break(doc)
    add_title(doc, 'Briefing e parâmetros de execução', 'Questionários vinculados às tarefas e dados para operações com IA ou execução híbrida.')
    heading(doc, '11 Questionário de briefing', 1)
    simple_form(doc, [
        ('Nome do questionário', '____________________________________________________________'),
        ('Descrição e objetivo', '____________________________________________________________'),
        ('Vincular às tarefas', '____________________________________________________________'),
    ])
    matrix(doc, ['Ordem', 'Pergunta para o cliente', 'Tipo de resposta', 'Obrigatória?', 'Ajuda ou exemplo'], [
        ['____', '________________________________________________', '☐ Texto ☐ Número ☐ Escolha ☐ Arquivo', '☐ Sim ☐ Não', '____________________________'],
        ['____', '________________________________________________', '☐ Texto ☐ Número ☐ Escolha ☐ Arquivo', '☐ Sim ☐ Não', '____________________________'],
        ['____', '________________________________________________', '☐ Texto ☐ Número ☐ Escolha ☐ Arquivo', '☐ Sim ☐ Não', '____________________________'],
        ['____', '________________________________________________', '☐ Texto ☐ Número ☐ Escolha ☐ Arquivo', '☐ Sim ☐ Não', '____________________________'],
        ['____', '________________________________________________', '☐ Texto ☐ Número ☐ Escolha ☐ Arquivo', '☐ Sim ☐ Não', '____________________________'],
    ], [.55,3.1,1.45,.85,1.3], 8)
    heading(doc, '12 Parâmetros para tarefa IA ou híbrida', 1)
    note(doc, 'Preencher somente para tarefas com modo IA ou híbrido.')
    matrix(doc, ['Tarefa', 'Provedor e modelo', 'Tokens entrada', 'Tokens saída', 'Custo por 1 mil entrada', 'Custo por 1 mil saída', 'Rodadas revisão', 'Revisão humana?'], [
        ['____________', '________________', '________', '________', 'R$ ________', 'R$ ________', '______', '☐ Sim ☐ Não'],
        ['____________', '________________', '________', '________', 'R$ ________', 'R$ ________', '______', '☐ Sim ☐ Não'],
    ], [.8,1.1,.8,.75,1.1,1.1,.9,1.0], 7.2)
    line(doc, 'Observação de custo, limite ou qualidade da IA', 3)

    page_break(doc)
    add_title(doc, 'Prazo custos e preço comercial', 'Campos que determinam a simulação. O sistema calcula preço e prazo no servidor.')
    heading(doc, '13 Prazo comercial', 1)
    simple_form(doc, [
        ('Prazo comercial base', '________ dias'),
        ('Prazo é definitivo?', '☐ Sim   ☐ Não - pendência comercial a resolver antes da publicação'),
        ('Motivo / premissas do prazo', '____________________________________________________________'),
    ])
    heading(doc, '14 Valor hora por especialidade', 1)
    note(doc, 'Hoje o valor/hora é configurado por especialidade e aplicado automaticamente a todos os produtos que usam aquela especialidade. Preencha aqui a decisão ou a referência que o administrador deve conferir no módulo global de precificação.')
    matrix(doc, ['Especialidade', 'Valor hora vigente ou proposto', 'Já está salvo na plataforma?', 'Responsável pela decisão', 'Observação'], [
        ['____________________', 'R$ ______________ / hora', '☐ Sim ☐ Não', '____________________', '____________________________'],
        ['____________________', 'R$ ______________ / hora', '☐ Sim ☐ Não', '____________________', '____________________________'],
        ['____________________', 'R$ ______________ / hora', '☐ Sim ☐ Não', '____________________', '____________________________'],
        ['____________________', 'R$ ______________ / hora', '☐ Sim ☐ Não', '____________________', '____________________________'],
    ], [1.5,1.4,1.25,1.35,1.4], 8)
    heading(doc, '15 Taxas e margem da operação', 1)
    note(doc, 'Esses parâmetros também são globais no módulo de precificação. Use este quadro para registrar a decisão e verificar se o simulador está usando a configuração correta.')
    matrix(doc, ['Parâmetro', 'Percentual ou valor', 'Confirmado por', 'Data'], [
        ['Imposto', '________ %', '________________________', '____ / ____ / ______'],
        ['Comissão', '________ %', '________________________', '____ / ____ / ______'],
        ['Taxa operacional', '________ %', '________________________', '____ / ____ / ______'],
        ['Margem de lucro', '________ %', '________________________', '____ / ____ / ______'],
        ['Reserva para revisão humana', '________ %', '________________________', '____ / ____ / ______'],
    ], [2.1,1.45,2.1,1.1])
    heading(doc, '16 Resultado esperado da simulação', 1)
    matrix(doc, ['Cenário', 'Variação e adicionais selecionados', 'Quantidade', 'Custo direto', 'Preço mínimo', 'Preço comercial final', 'Prazo resultante'], [
        ['Base', '________________________________', '____', 'R$ ______', 'R$ ______', 'R$ ______', '______ dias'],
        ['Alternativo 1', '________________________________', '____', 'R$ ______', 'R$ ______', 'R$ ______', '______ dias'],
        ['Alternativo 2', '________________________________', '____', 'R$ ______', 'R$ ______', 'R$ ______', '______ dias'],
    ], [.8,2.0,.7,1.0,1.0,1.2,1.0], 7.8)

    page_break(doc)
    add_title(doc, 'Revisão teste e publicação', 'Use esta página somente depois de concluir as seções anteriores e cadastrar o produto localmente.')
    heading(doc, '17 Checklist de produto 100 por cento cadastrado', 1)
    checklist = [
        'Nome interno, slug único, título comercial, imagem e textos preenchidos.',
        'Pilar, categoria e classificações 4F definidos.',
        'Variações, opções, adicionais e condições cadastrados ou marcados como não aplicáveis.',
        'Todas as tarefas têm nome, especialidade, modo de execução e minutos estimados definidos.',
        'Etapas, dependências, questionários e perguntas obrigatórias revisados.',
        'Prazo comercial base definido; nenhum prazo provisório usado como decisão final.',
        'Valores/hora das especialidades e taxas globais conferidos no módulo de precificação.',
        'Simulação base e cenários de variação conferidos: custo, preço comercial e prazo coerentes.',
        'Se mensal recorrente, recorrência marcada e período mensal configurado com desconto explícito quando aplicável.',
        'Pré-visualização administrativa revisada e catálogo do cliente testado em ambiente local.',
        'Resumo de alteração preenchido e versão publicada somente após as pendências estruturais zerarem.',
        'Status alterado para Disponível apenas depois de existir versão publicada e produto estar ofertável.',
    ]
    for item in checklist:
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(5)
        r = p.add_run('☐  ' + item); r.font.size = Pt(10.2)
    heading(doc, '18 Evidência dos testes', 1)
    matrix(doc, ['Teste', 'Resultado esperado', 'Resultado observado', 'Responsável e data'], [
        ['Cadastro de informações', 'Salva e confirma sucesso', '________________________________', '____________________________'],
        ['Classificação', 'Aparece nos filtros certos', '________________________________', '____________________________'],
        ['Variação e adicional', 'Altera preço, prazo ou entrega conforme regra', '________________________________', '____________________________'],
        ['Briefing', 'Perguntas corretas aparecem na tarefa', '________________________________', '____________________________'],
        ['Simulação', 'Custo, preço e prazo coerentes', '________________________________', '____________________________'],
        ['Publicação e catálogo cliente', 'Produto aparece conforme status', '________________________________', '____________________________'],
    ], [1.25,1.7,2.2,1.45], 8)
    heading(doc, '19 Aprovação para cadastrar e publicar', 1)
    simple_form(doc, [
        ('Responsável pelo cadastro', '____________________________________________________________'),
        ('Responsável comercial', '____________________________________________________________'),
        ('Aprovação da operação', '☐ Aprovado   ☐ Aprovado com observações   ☐ Não aprovado'),
        ('Observações finais', '____________________________________________________________'),
        ('Versão publicada e data', '____________________________________________________________'),
    ])

    doc.save(OUT)
    print(OUT)

if __name__ == '__main__':
    build()
